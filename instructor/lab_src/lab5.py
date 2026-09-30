TITLE = "Automate the Monthly Report"
SLUG = "automation"

CELLS = [
("md", """
**Scenario.** Every month, someone at Limestone Outfitters spends a day building the same Excel sales report by hand: they download exchange rates, copy data, build pivots and format the tables. Your job is to turn that day into **one function call**.

By the end you will have:
1. pulled live data from a web **API**,
2. wrapped your analysis in a **function** with a parameter (`month`),
3. produced a formatted **Excel report** a manager would actually open,
4. generated six months of reports in a **loop**, and checked that they reconcile.
"""),
("code", """
import pandas as pd, requests, os
orders = pd.read_csv(DATA + "orders.csv", parse_dates=["order_date"])
items = pd.read_csv(DATA + "order_items.csv")
products = pd.read_csv(DATA + "products.csv")
customers = pd.read_csv(DATA + "customers.csv")
returns = pd.read_csv(DATA + "returns.csv", parse_dates=["return_date"])
"""),
("md", """
## Part 1: Call an API (10 min)
An API is a URL that returns **data** instead of a web page. The European Central Bank publishes daily exchange rates. Frankfurter serves them for free, with no key needed.

Open this in a new browser tab first and look at what comes back:
`https://api.frankfurter.dev/v1/2026-03-01..2026-03-31?base=USD&symbols=EUR,GBP`

That format is **JSON**: nested `{key: value}` pairs. Python reads it into a dictionary.
"""),
("code", """
url = "https://api.frankfurter.dev/v1/2026-03-01..2026-03-31"
resp = requests.get(url, params={"base": "USD", "symbols": "EUR,GBP"}, timeout=20)
print(resp.status_code)          # 200 = OK
data = resp.json()
print(list(data.keys()))
list(data["rates"].items())[:3]
"""),
("md", """
**1a.** Write a function `fx_for_month(month)` that takes `"2026-03"` and returns a dict like `{"USD": 1.0, "EUR": 1.08, "GBP": 1.27}`: **USD per 1 unit** of each currency, averaged over the month.

Notice that the API gives *EUR per USD*, and you need the reverse. It is also good practice to fall back to the saved file `fx_rates_monthly.csv` if the API is down, because automation that breaks when a website hiccups isn't automation.
"""),
("code", """
def fx_for_month(month):
    # your code (or AI's code, which you have read line by line)
    ...
fx_for_month("2026-03")
""", """
def fx_for_month(month):
    \"\"\"USD per 1 unit of each currency, monthly average. API first, saved file as fallback.\"\"\"
    start = pd.Period(month).start_time.date()
    end = pd.Period(month).end_time.date()
    try:
        r = requests.get(f"https://api.frankfurter.dev/v1/{start}..{end}",
                         params={"base": "USD", "symbols": "EUR,GBP"}, timeout=20)
        r.raise_for_status()
        per_usd = pd.DataFrame(r.json()["rates"]).T.mean()          # EUR per USD, GBP per USD
        source = "api"
    except Exception as e:
        saved = pd.read_csv(DATA + "fx_rates_monthly.csv")
        per_usd = saved[saved.month == month].set_index("currency").per_usd.drop("USD")
        source = f"saved file ({type(e).__name__})"
    rates = {"USD": 1.0, **{c: 1 / v for c, v in per_usd.items()}}
    print(f"FX {month} from {source}: " + ", ".join(f"{c} {v:.4f}" for c, v in rates.items()))
    return rates
fx_for_month("2026-03")
"""),
("md", """
## Part 2: From analysis to a function (15 min)
**2a.** Write `monthly_kpis(month)` that returns a dict of DataFrames:
- `summary`: one column of KPIs: net revenue (USD), completed orders, average order value (USD), units sold, return rate (units returned ÷ units sold for this month's orders), # new customers (first order this month)
- `by_category`: net revenue USD and units by category, sorted
- `by_country`: net revenue USD and orders by country, sorted

Rules (same company definitions as Lab 1): completed orders only, no test accounts, net revenue = qty × price × (1 − discount), converted with `fx_for_month`.

**Tip:** get it working for one month *in plain cells* first, then indent it into a function. Ask AI to help, but give it these rules.
"""),
("code", """
def monthly_kpis(month):
    ...
k = monthly_kpis("2026-03")
k["summary"]
""", """
real = set(customers.loc[~customers.is_test, "customer_id"])
first_order = (orders[(orders.status == "Completed") & orders.customer_id.isin(real)].groupby("customer_id").order_date.min().dt.strftime("%Y-%m"))

def monthly_kpis(month):
    fx = fx_for_month(month)
    o = orders[(orders.order_date.dt.strftime("%Y-%m") == month) & (orders.status == "Completed")
               & orders.customer_id.isin(real)]
    li = o.merge(items, on="order_id", validate="one_to_many").merge(products, on="product_id", validate="many_to_one")
    li["net_usd"] = li.quantity * li.unit_price * (1 - li.discount_pct) * li.currency.map(fx)
    ret_units = (returns.merge(li[["order_id", "product_id"]].drop_duplicates(), on=["order_id", "product_id"])
                        .qty_returned.sum())
    summary = pd.DataFrame({"value": {
        "Net revenue (USD)": li.net_usd.sum(),
        "Completed orders": o.order_id.nunique(),
        "Average order value (USD)": li.net_usd.sum() / o.order_id.nunique(),
        "Units sold": li.quantity.sum(),
        "Return rate (units, this month's orders)": ret_units / li.quantity.sum(),
        "New customers": int((first_order == month).sum()),
    }})
    by_cat = (li.groupby("category").agg(net_revenue_usd=("net_usd", "sum"), units=("quantity", "sum"))
                .sort_values("net_revenue_usd", ascending=False).reset_index())
    by_cty = (li.groupby("country").agg(net_revenue_usd=("net_usd", "sum"), orders=("order_id", "nunique"))
                .sort_values("net_revenue_usd", ascending=False).reset_index())
    return {"summary": summary, "by_category": by_cat, "by_country": by_cty, "month": month}

k = monthly_kpis("2026-03")
k["summary"]
"""),
("md", """
## Part 3: Write an Excel report someone would open (15 min)
**3a.** Write `build_report(month, folder="reports")` that saves `reports/sales_2026-03.xlsx` with three sheets: *Summary*, *By Category*, *By Country*. Make it look finished:
- bold header row, sensible column widths
- currency format `$#,##0` for money, `0.0%` for rates
- a bar chart of revenue by category on the *By Category* sheet
- a note at the bottom of *Summary*: data source, FX source, and when the report was generated

`openpyxl` does all of this. The AI knows it well, but check every number format it chooses.
"""),
("code", """
def build_report(month, folder="reports"):
    ...
build_report("2026-03")
""", """
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font
from datetime import datetime

def build_report(month, folder="reports"):
    k = monthly_kpis(month)
    os.makedirs(folder, exist_ok=True)
    path = f"{folder}/sales_{month}.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        k["summary"].rename_axis("KPI").to_excel(xw, sheet_name="Summary")
        k["by_category"].to_excel(xw, sheet_name="By Category", index=False)
        k["by_country"].to_excel(xw, sheet_name="By Country", index=False)
        wb = xw.book
        for ws in wb.worksheets:
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = max(14, max(len(str(c.value or "")) for c in col) + 2)
        s = wb["Summary"]
        fmt = {"Net revenue (USD)": "$#,##0", "Average order value (USD)": "$#,##0.00",
               "Return rate (units, this month's orders)": "0.0%"}
        for row in s.iter_rows(min_row=2, max_col=2):
            row[1].number_format = fmt.get(row[0].value, "#,##0")
        s.cell(row=s.max_row + 2, column=1,
               value=f"Source: order system extract; FX: ECB via frankfurter.dev (monthly avg). Generated {datetime.now():%Y-%m-%d %H:%M}.")
        for name in ["By Category", "By Country"]:
            for c in wb[name]["B"][1:]:
                c.number_format = "$#,##0"
        ws = wb["By Category"]
        chart = BarChart(); chart.title = f"Net revenue by category, {month}"; chart.y_axis.title = "USD"
        chart.add_data(Reference(ws, min_col=2, min_row=1, max_row=ws.max_row), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=2, max_row=ws.max_row))
        chart.legend = None
        ws.add_chart(chart, "E2")
    print("saved", path)
    return k

_ = build_report("2026-03")
"""),
("md", """
## Part 4: Run it for six months, and prove it reconciles (10 min)
**4a.** Loop over `2026-01` … `2026-06` and build all six reports.

**4b. Reconciliation check.** Recompute H1 2026 net revenue **a different way**: one big query over all six months at once, using `fx_rates_monthly.csv` instead of the API. The sum of your six reports must match within 0.5%. (Why might it not match *exactly*? Write your answer.)
"""),
("code", """
# 4a
# 4b
""", """
monthly = {}
for m in pd.period_range("2026-01", "2026-06", freq="M").astype(str):
    monthly[m] = build_report(m)["summary"].loc["Net revenue (USD)", "value"]

fx = pd.read_csv(DATA + "fx_rates_monthly.csv")
o = orders[(orders.order_date >= "2026-01-01") & (orders.order_date < "2026-07-01") & (orders.status == "Completed")
           & orders.customer_id.isin(real)].assign(month=lambda d: d.order_date.dt.strftime("%Y-%m"))
li = o.merge(items, on="order_id").merge(fx, on=["month", "currency"])
independent = (li.quantity * li.unit_price * (1 - li.discount_pct) * li.usd_per_unit).sum()
reports_total = sum(monthly.values())
print(f"sum of reports {reports_total:,.0f} | independent {independent:,.0f} | diff {reports_total / independent - 1:.3%}")
assert abs(reports_total / independent - 1) < 0.005
"""),
("md", """
## Part 5: Make it run without you (discussion + homework)
A function is only half of automation. Something still has to **call** it on the first of every month. Options, from least to most engineering:

| Option | How | Good for |
|---|---|---|
| Colab / Jupyter scheduled run | Colab Pro scheduling or a laptop cron job | personal use |
| **GitHub Actions** | a YAML file that runs your script on a cron schedule and emails or uploads the file | small teams, free |
| Power Automate / Zapier | trigger → run a script → post to Teams/Slack/SharePoint | Microsoft-shop companies |
| Airflow / dbt / cloud schedulers | pipelines owned by a data team | production |

**Homework:** ask an AI to write a GitHub Actions workflow that runs `build_report` for the *previous month* at 7am on the 1st of every month. Before you trust it, explain **every line** of the YAML in a comment, and name one way it could fail silently.
"""),
]
